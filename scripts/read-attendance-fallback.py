#!/usr/bin/env python3
"""Reads the manually-maintained "Lofty TSR Attendance" SharePoint/OneDrive workbook and dumps
{employeeName: {"YYYY-MM-DD": "CODE"}} as JSON on stdout - used as a fallback source when the
portal's own attendance entry is missing for that employee/date. Employee names must match the
roster's employeeName field exactly (confirmed live: they do, verbatim, including quirks like
'Reymark "Mac" Lopez' and 'CristineJoy Banacia').

Sheet layout: one tab per month ("Sep 2026"), row 1 = day-of-month numbers starting at column C,
row 2 = weekday abbreviations, rows 3+ = one row per employee (Employee, Team Lead, then one
status code per day column). Status codes already match the portal's own vocabulary (ONSITE,
WFH, RD, NCNS, SL, PTO, LATE, EL, SL-HD, EL-HD, A) - confirmed live via a full-sheet code scan.
"""
import json
import sys

import openpyxl

DEFAULT_PATH = '/Users/mac/Library/CloudStorage/OneDrive-MoatableInc/Lofty PH Repository - Documents/Lofty TSR/Attendance/2026.xlsx'

MONTHS = {
    'Jan': 1, 'Feb': 2, 'Mar': 3, 'Apr': 4, 'May': 5, 'Jun': 6,
    'Jul': 7, 'Aug': 8, 'Sep': 9, 'Oct': 10, 'Nov': 11, 'Dec': 12,
}


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PATH
    wb = openpyxl.load_workbook(path, data_only=True)
    result = {}
    for sheet_name in wb.sheetnames:
        parts = sheet_name.split()
        if len(parts) != 2 or parts[0] not in MONTHS:
            continue
        month = MONTHS[parts[0]]
        try:
            year = int(parts[1])
        except ValueError:
            continue
        ws = wb[sheet_name]
        day_cols = {}
        for col in range(3, ws.max_column + 1):
            val = ws.cell(row=1, column=col).value
            if isinstance(val, (int, float)) and float(val).is_integer() and 1 <= val <= 31:
                day_cols[col] = int(val)
        for r in range(3, ws.max_row + 1):
            name = ws.cell(row=r, column=1).value
            if not name:
                continue
            name = str(name).strip()
            for col, day in day_cols.items():
                code = ws.cell(row=r, column=col).value
                if code is None or code == '':
                    continue
                code = str(code).strip()
                if not code:
                    continue
                date_str = f"{year:04d}-{month:02d}-{day:02d}"
                result.setdefault(name, {})[date_str] = code
    json.dump(result, sys.stdout)


if __name__ == '__main__':
    main()
